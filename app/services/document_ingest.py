import asyncio
from io import BytesIO
import logging
import uuid
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from app.services.mongo_checkpoint import doc_vector_store


logger = logging.getLogger(__name__)

text_splitter = RecursiveCharacterTextSplitter(
    chunk_size = 1000,
    chunk_overlap=200,
    add_start_index = True
)

def extract_text(file_bytes: bytes, file_name: str, mime_type: str) -> str:
    """Extract raw text depending on file type."""
    if mime_type == "application/pdf" or file_name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    elif file_name.endswith(".docx"):
        from docx import Document as DocxDocument
        doc = DocxDocument(BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs)

    elif mime_type.startswith("text/") or file_name.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore")

    else:
        raise ValueError(f"Unsupported file type: {mime_type} / {file_name}")


def _ingest_document(file_bytes: bytes, file_name: str, mime_type: str, user_id: str) -> int:
    """Full pipeline: extract -> chunk -> embed -> store. Returns number of chunks stored."""
    document_id = str(uuid.uuid4())
    text = extract_text(file_bytes, file_name, mime_type)

    if not text.strip():
        raise ValueError("No extractable text found in file.")

    docs = [
        Document(
            page_content=text,
            metadata={
                "user_id": str(user_id),
                "source": file_name,
                "document_id":document_id
            }
        )
    ]

    all_splits = text_splitter.split_documents(docs)
    logger.info(f"Split '{file_name}' into {len(all_splits)} chunks")

    doc_vector_store.add_documents(all_splits)
    logger.info(f"Ingested {len(all_splits)} chunks from '{file_name}' for user {user_id}")

    return len(all_splits)

async def ingest_document(file_bytes: bytes, file_name: str, mime_type: str, user_id: str) -> int:
    """Async entry point: extract -> chunk -> embed -> store.
 
    Runs the full pipeline in a background thread (asyncio.to_thread) so
    a large PDF's parsing + embedding calls don't freeze the event loop —
    and every other Telegram user's messages — while it runs.
 
    NOTE: this only stops it from blocking the *server* for everyone
    else. The uploader themselves still waits the same wall-clock time
    for their own reply, since telegram.py currently awaits this call
    before it can continue building the message. Making it feel instant
    for the uploader (e.g. an immediate "processing your document..."
    reply, then a follow-up message once ingestion finishes) needs a
    small change in telegram.py/routes.py — happy to do that next.
    """
    return await asyncio.to_thread(
        _ingest_document, file_bytes, file_name, mime_type, user_id
    )
 